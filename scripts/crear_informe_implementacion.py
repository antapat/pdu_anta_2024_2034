from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "documentacion"
ASSET_DIR = OUT_DIR / "_informe_assets"
OUTPUT = OUT_DIR / "Informe_Implementacion_GeoVisor_PDU_Anta.docx"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)

GREEN = "1B4D3E"
GREEN_DARK = "0F2D24"
GREEN_PALE = "#EBF3EF"
GOLD = "C59B27"
GOLD_PALE = "FDF9EE"
SLATE = "475569"
LIGHT = "F4F6F5"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    set_font(run, size=7.5, color=SLATE)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_picture_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    set_paragraph_spacing(p, after=3)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    set_paragraph_spacing(p, after=3)
    return p


def add_manual_number(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.add_run(f"{number}.\t{text}")
    set_paragraph_spacing(p, after=3)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=7, line=1.14)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, size=9, italic=True, color=SLATE)
    set_paragraph_spacing(p, before=2, after=8, line=1.0)
    return p


def add_table(doc, headers, rows, widths=None, small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, GREEN)
        set_cell_border(cell)
        set_cell_margins(cell, 115, 120, 115, 120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        set_font(r, size=8.5 if small else 9, bold=True, color=WHITE)
        set_paragraph_spacing(p, after=0, line=1.0)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell, 105, 120, 105, 120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=8.3 if small else 8.8, color=BLACK)
            set_paragraph_spacing(p, after=0, line=1.03)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (key, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Cm(5.0)
        cells[1].width = Cm(11.5)
        set_cell_shading(cells[0], GREEN_PALE)
        set_cell_shading(cells[1], WHITE)
        for cell in cells:
            set_cell_border(cell)
            set_cell_margins(cell, 115, 140, 115, 140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        r0 = cells[0].paragraphs[0].add_run(key)
        set_font(r0, size=9, bold=True, color=GREEN_DARK)
        r1 = cells[1].paragraphs[0].add_run(value)
        set_font(r1, size=9, color=BLACK)
        set_paragraph_spacing(cells[0].paragraphs[0], after=0)
        set_paragraph_spacing(cells[1].paragraphs[0], after=0)
    return table


def build_architecture_diagram(path):
    width, height = 1800, 920
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = "C:/Windows/Fonts/arial.ttf"
    bold_path = "C:/Windows/Fonts/arialbd.ttf"
    title_font = ImageFont.truetype(bold_path, 42)
    head_font = ImageFont.truetype(bold_path, 28)
    body_font = ImageFont.truetype(font_path, 23)
    small_font = ImageFont.truetype(font_path, 20)

    draw.text((width // 2, 44), "Arquitectura operativa del GeoVisor PDU Anta", fill="#000000", font=title_font, anchor="ma")

    boxes = {
        "citizen": (80, 230, 390, 475, "Población usuaria", "Consulta punto o predio\nVisualiza situación PDU\nGenera reporte informativo", GREEN_PALE),
        "shiny": (560, 180, 1010, 530, "Aplicación Shiny", "Interfaz bslib y Leaflet\nAnálisis espacial con sf\nReportes Quarto y Typst\nEjecución en Posit Connect", "#FDF9EE"),
        "supabase": (1180, 150, 1715, 560, "Supabase", "PostgreSQL y PostGIS\n8 capas espaciales\nParámetros urbanísticos\nMétricas y auditoría\nRLS y funciones RPC", "#EEF4F8"),
        "admin": (235, 670, 675, 850, "Carga administrativa", "Migraciones SQL y carga transaccional desde datos oficiales", "#F4F6F5"),
        "ops": (1050, 680, 1550, 850, "Operación institucional", "Verificación de conteos, índices GiST, secretos y monitoreo", "#F4F6F5"),
    }
    for key, (x1, y1, x2, y2, head, body, fill) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=fill, outline="#1B4D3E", width=4)
        draw.text(((x1 + x2) // 2, y1 + 32), head, fill="#000000", font=head_font, anchor="ma")
        y = y1 + 92
        for line in body.split("\n"):
            draw.text(((x1 + x2) // 2, y), line, fill="#263238", font=body_font, anchor="ma")
            y += 42

    def arrow(start, end, label):
        draw.line((start, end), fill="#1B4D3E", width=8)
        ex, ey = end
        sx, sy = start
        import math
        angle = math.atan2(ey - sy, ex - sx)
        size = 22
        points = [
            (ex, ey),
            (ex - size * math.cos(angle - 0.55), ey - size * math.sin(angle - 0.55)),
            (ex - size * math.cos(angle + 0.55), ey - size * math.sin(angle + 0.55)),
        ]
        draw.polygon(points, fill="#1B4D3E")
        mx, my = (sx + ex) // 2, (sy + ey) // 2
        draw.rounded_rectangle((mx - 105, my - 24, mx + 105, my + 24), radius=8, fill="white")
        draw.text((mx, my), label, fill="#1B4D3E", font=small_font, anchor="mm")

    arrow((390, 350), (560, 350), "HTTPS")
    arrow((1010, 350), (1180, 350), "PostgreSQL SSL")
    arrow((675, 760), (1180, 515), "Carga controlada")
    arrow((1300, 680), (1420, 560), "Validación")
    img.save(path, quality=95)


doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.14

for name, size, before, after in (("Title", 25, 0, 10), ("Heading 1", 15, 16, 8), ("Heading 2", 12, 12, 6), ("Heading 3", 10.5, 9, 4)):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

styles["Caption"].font.name = "Arial"
styles["Caption"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Caption"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

# Header and footer
header = section.header
header.distance = Cm(0.7)
h_table = header.add_table(rows=1, cols=3, width=Cm(16.4))
h_table.alignment = WD_TABLE_ALIGNMENT.CENTER
h_table.columns[0].width = Cm(1.5)
h_table.columns[1].width = Cm(12.8)
h_table.columns[2].width = Cm(2.1)
shield = ROOT / "www" / "escudo_muni_anta.png"
pdu_logo = ROOT / "www" / "pdu_actualizado_sf.png"
if shield.exists():
    pic = h_table.cell(0, 0).paragraphs[0].add_run().add_picture(str(shield), height=Cm(0.8))
    set_picture_alt(pic, "Escudo de la Municipalidad Provincial de Anta")
hp = h_table.cell(0, 1).paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run("MUNICIPALIDAD PROVINCIAL DE ANTA\nINFORME DE IMPLEMENTACIÓN DEL GEOVISOR PDU")
set_font(hr, size=7.5, bold=True, color=SLATE)
if pdu_logo.exists():
    h_table.cell(0, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pic = h_table.cell(0, 2).paragraphs[0].add_run().add_picture(str(pdu_logo), height=Cm(0.75))
    set_picture_alt(pic, "Logotipo del Plan de Desarrollo Urbano de Anta 2024 2034")
for row in h_table.rows:
    for cell in row.cells:
        set_cell_margins(cell, 20, 30, 20, 30)

footer = section.footer
footer.distance = Cm(0.7)
f_table = footer.add_table(rows=1, cols=3, width=Cm(16.4))
f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
f_table.columns[0].width = Cm(6.0)
f_table.columns[1].width = Cm(7.5)
f_table.columns[2].width = Cm(2.9)
for idx, text in enumerate(("Municipalidad Provincial de Anta", "GeoVisor PDU Anta 2024 2034", "Página ")):
    p = f_table.cell(0, idx).paragraphs[0]
    p.alignment = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT][idx]
    r = p.add_run(text)
    set_font(r, size=7.5, color=SLATE)
    if idx == 2:
        add_field(p, "PAGE")

# Cover
cover_table = doc.add_table(rows=1, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_table.autofit = False
cover_table.columns[0].width = Cm(8.2)
cover_table.columns[1].width = Cm(8.2)
if shield.exists():
    p = cover_table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pic = p.add_run().add_picture(str(shield), height=Cm(2.6))
    set_picture_alt(pic, "Escudo de la Municipalidad Provincial de Anta")
if pdu_logo.exists():
    p = cover_table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pic = p.add_run().add_picture(str(pdu_logo), height=Cm(2.4))
    set_picture_alt(pic, "Logotipo del Plan de Desarrollo Urbano de Anta 2024 2034")

doc.add_paragraph().paragraph_format.space_after = Pt(20)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Informe de implementación del GeoVisor PDU Anta 2024 2034")
set_font(r, size=25, bold=True, color=BLACK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Arquitectura Supabase PostGIS aplicación Shiny y evaluación de interfaz")
set_font(r, size=14, bold=True, color=GREEN_DARK)
set_paragraph_spacing(p, after=22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Proyecto")
set_font(r, size=10, bold=True, color=SLATE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Mejoramiento de los Servicios de Gestión Territorial y Desarrollo Urbano Sostenible del distrito de Anta provincia de Anta departamento del Cusco")
set_font(r, size=12, bold=True, color=BLACK)
set_paragraph_spacing(p, after=26, line=1.1)

cover_meta = add_key_value_table(doc, [
    ("Entidad", "Municipalidad Provincial de Anta"),
    ("Unidad responsable", "Gerencia de Desarrollo Urbano y Rural"),
    ("Producto", "GeoVisor informativo del Plan de Desarrollo Urbano de Anta 2024 2034"),
    ("Versión del informe", "1.0"),
    ("Fecha", "9 de octubre de 2026"),
    ("Lugar", "Anta Cusco"),
])

doc.add_paragraph().paragraph_format.space_after = Pt(12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Documento técnico institucional")
set_font(r, size=10, italic=True, color=SLATE)

doc.add_page_break()

doc.add_heading("Control documental", level=1)
add_key_value_table(doc, [
    ("Título", "Informe de implementación del GeoVisor PDU Anta 2024 2034"),
    ("Finalidad", "Documentar la solución implementada, su estado operativo y las mejoras recomendadas para la interfaz pública"),
    ("Ámbito", "Arquitectura, información espacial, consultas, reportes, seguridad, despliegue, validación y experiencia de usuario"),
    ("Carácter", "Informe técnico del producto digital del proyecto"),
    ("Estado", "Implementación de código completada. Carga remota y publicación sujetas a la configuración de credenciales institucionales"),
])

doc.add_heading("Contenido", level=1)
contents = [
    "Resumen ejecutivo", "Antecedentes del proyecto", "Objeto y alcance del informe",
    "Descripción del GeoVisor", "Arquitectura de la solución", "Información espacial y modelo de datos",
    "Implementación técnica", "Seguridad y protección de datos", "Despliegue y operación",
    "Validación realizada", "Estado de implementación", "Evaluación de la interfaz",
    "Propuesta de mejora", "Hoja de ruta", "Riesgos y medidas de control", "Conclusiones y recomendaciones",
    "Anexos técnicos",
]
for item in contents:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(item)
    set_font(r, size=9.5, color=SLATE)
    set_paragraph_spacing(p, after=2)

doc.add_page_break()

doc.add_heading("Resumen ejecutivo", level=1)
add_body(doc, "El GeoVisor PDU Anta es el componente digital de consulta territorial del proyecto de mejoramiento de los servicios de gestión territorial y desarrollo urbano sostenible. Su finalidad es permitir que la población identifique la situación de un punto o predio respecto del Plan de Desarrollo Urbano de Anta 2024 2034, consulte la zonificación, las vías propuestas, las servidumbres y los parámetros urbanísticos referenciales, y genere un reporte informativo.")
add_body(doc, "La solución combina una aplicación Shiny publicada en Posit Connect con una base de datos PostgreSQL y PostGIS en Supabase. La arquitectura se modificó para que Supabase sea la fuente espacial oficial. La aplicación ya no utiliza archivos RDS durante la ejecución y se detiene si la conexión institucional o las tablas requeridas no están disponibles. Los RDS se conservan solamente como insumo administrativo para la carga controlada.")
add_body(doc, "La implementación incorpora ocho capas espaciales, una matriz de cuarenta registros de parámetros urbanísticos, consultas por punto y polígono, carga de archivos prediales, generación de reportes PDF, registro de consultas, contador de visitas y contador de constancias. La opción pública de envío de aportes fue retirada porque el propósito del servicio es informativo; además, se revocó la inserción anónima en la tabla histórica de reportes ciudadanos.")
add_body(doc, "La revisión de interfaz concluye que la aplicación cuenta con identidad institucional consistente y funciones territoriales completas, pero necesita simplificar la cabecera, reducir elementos decorativos, mejorar la experiencia móvil, reforzar la accesibilidad y ordenar el flujo de consulta. Las mejoras prioritarias son completar la carga remota en Supabase, eliminar el código residual del antiguo buzón, mostrar la fecha de actualización de los datos y validar el recorrido completo con usuarios de la Municipalidad y población del ámbito del PDU.")

doc.add_heading("Antecedentes del proyecto", level=1)
add_body(doc, "El producto se desarrolla en el marco del proyecto denominado Mejoramiento de los Servicios de Gestión Territorial y Desarrollo Urbano Sostenible del distrito de Anta provincia de Anta departamento del Cusco. El instrumento territorial consultado es el Plan de Desarrollo Urbano de Anta 2024 2034, cuya fecha de aprobación consignada en la aplicación es el 4 de abril de 2025.")
add_body(doc, "La necesidad institucional consiste en ofrecer a la población una consulta accesible de la información técnica del PDU sin reemplazar los procedimientos formales de la Municipalidad. Por esa razón, el sistema presenta resultados referenciales y advierte que el reporte generado no constituye certificado oficial ni genera derechos urbanísticos.")

doc.add_heading("Objeto y alcance del informe", level=1)
add_body(doc, "Este informe registra las decisiones técnicas y funcionales del GeoVisor, describe la transición de archivos locales a Supabase PostGIS, presenta las medidas de seguridad y operación, documenta las validaciones realizadas y formula una propuesta de mejora de la interfaz institucional.")
add_body(doc, "El alcance comprende el código y los recursos presentes en el repositorio del producto. La evaluación de interfaz se basa en la estructura Shiny, los estilos CSS, los controles implementados y la salida cartográfica disponible. La verificación en producción deberá completarse después de cargar las tablas en el proyecto Supabase y configurar los secretos en Posit Connect.")

doc.add_heading("Descripción del GeoVisor", level=1)
add_body(doc, "El GeoVisor organiza la consulta pública en cuatro áreas visibles: visor territorial, reporte informativo, catálogo normativo de zonas e información institucional del PDU. El antiguo buzón ciudadano se deshabilitó para mantener el servicio dentro de su objetivo informativo.")
add_table(doc, ["Componente", "Función para la población", "Resultado"], [
    ("Visor territorial", "Consulta un punto mediante coordenadas UTM o geográficas, GPS móvil o un polígono predial", "Ficha territorial y representación en mapa"),
    ("Reporte informativo", "Consolida los resultados de la consulta activa", "Vista institucional y descarga PDF cuando el motor está disponible"),
    ("Catálogo normativo", "Busca códigos, nombres y usos permitidos", "Parámetros urbanísticos por subzona"),
    ("Información del PDU", "Explica el instrumento y su alcance", "Contexto institucional y advertencia legal"),
    ("Métricas públicas", "Muestra el uso general del servicio", "Visitas aproximadas y constancias generadas"),
], widths=[3.5, 7.5, 5.5])

doc.add_heading("Arquitectura de la solución", level=1)
diagram_path = ASSET_DIR / "arquitectura_geovisor.png"
build_architecture_diagram(diagram_path)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic = p.add_run().add_picture(str(diagram_path), width=Cm(16.0))
set_picture_alt(pic, "Diagrama de arquitectura del GeoVisor con población usuaria, aplicación Shiny, Supabase y operación institucional")
add_caption(doc, "Figura 1 Arquitectura operativa del GeoVisor PDU Anta")
add_body(doc, "La población accede a la aplicación Shiny mediante HTTPS. La aplicación se conecta a PostgreSQL en Supabase mediante SSL, carga las capas PostGIS y ejecuta el análisis espacial con el paquete sf. Las métricas pueden utilizar PostgreSQL o las funciones RPC de Supabase. Los procesos administrativos aplican migraciones, cargan los datos oficiales en una transacción y verifican los conteos antes de publicar.")
add_table(doc, ["Capa", "Tecnología", "Responsabilidad"], [
    ("Presentación", "R Shiny bslib Leaflet", "Interacción, mapas, formularios, fichas y navegación"),
    ("Procesamiento", "R sf dplyr ggplot2", "Intersecciones, distancias, áreas, mapas y preparación del reporte"),
    ("Persistencia", "Supabase PostgreSQL PostGIS", "Geometrías, parámetros, auditoría, métricas y trazabilidad de importaciones"),
    ("Reportes", "Quarto Typst y fallback pagedown", "Constancia informativa en PDF"),
    ("Alojamiento", "Posit Connect Cloud", "Ejecución de la aplicación y administración de secretos"),
    ("Versionamiento", "GitHub y migraciones SQL", "Control de código y evolución del esquema"),
], widths=[3.0, 4.5, 9.0])

doc.add_heading("Información espacial y modelo de datos", level=1)
add_body(doc, "El conjunto oficial se transforma a EPSG 4326 para la visualización web y conserva el uso de EPSG 32718 para cálculos métricos en el ámbito de Anta. La carga inicial prevista contiene 4501 geometrías y cuarenta registros de parámetros urbanísticos.")
add_table(doc, ["Tabla PostGIS", "Tipo", "Registros previstos", "Uso principal"], [
    ("ambito", "Polígono", "1", "Delimitación del ámbito de intervención"),
    ("distritos", "Polígono", "9", "Identificación distrital"),
    ("zonificacion", "Polígono", "2741", "Código de zona subzona y área normativa"),
    ("vias", "Línea", "1729", "Sistema vial propuesto y secciones"),
    ("alta_tension", "Línea", "7", "Infraestructura eléctrica"),
    ("alta_tension_buffer", "Polígono", "4", "Faja referencial de servidumbre"),
    ("ferrocarril", "Línea", "5", "Eje ferroviario"),
    ("ferrocarril_buffer", "Polígono", "5", "Franja referencial ferroviaria"),
    ("parametros_urbanisticos", "Tabla", "40", "Densidad, lote, altura, retiros, usos y validación"),
], widths=[4.0, 2.3, 3.0, 7.2], small=True)
add_body(doc, "Las tablas espaciales y la matriz urbanística tienen Row Level Security habilitada y no se exponen directamente a los roles anónimo o autenticado. La aplicación usa una conexión PostgreSQL institucional, cuyos valores deben almacenarse como secretos en el entorno de despliegue.")

doc.add_heading("Implementación técnica", level=1)
doc.add_heading("Preparación y carga de datos", level=2)
add_body(doc, "El proceso de preparación corrige geometrías, normaliza atributos, asigna colores y genera los insumos administrativos. El cargador transforma todas las capas a EPSG 4326, reemplaza las tablas previstas dentro de una transacción, crea índices GiST sobre cada geometría, habilita RLS, revoca el acceso público y registra el número de filas importadas.")
add_number(doc, "Aplicar las migraciones versionadas para habilitar PostGIS, auditoría y métricas.")
add_number(doc, "Cargar las ocho capas y la matriz urbanística con confirmación explícita del proyecto de destino.")
add_number(doc, "Verificar existencia, conteos, CRS y lectura de geometrías antes de desplegar Shiny.")

doc.add_heading("Consumo desde la aplicación Shiny", level=2)
add_body(doc, "Al iniciar, la aplicación exige una conexión PostgreSQL directa a Supabase. Comprueba la existencia de las nueve tablas requeridas, rechaza tablas vacías, lee cada capa con sf, valida el CRS y reconstruye las paletas de zonificación y vías. Si cualquiera de estas condiciones falla, el proceso se detiene con un mensaje de diagnóstico.")
add_body(doc, "Los archivos data pdu capas optimizadas y parametros urbanisticos no se incluyen en el manifiesto de Posit Connect. Esta decisión evita que el despliegue muestre información local cuando la fuente institucional no esté disponible.")

doc.add_heading("Consultas territoriales", level=2)
add_body(doc, "La consulta puntual acepta coordenadas UTM 18S, coordenadas geográficas y posición GPS. La consulta por polígono acepta vértices escritos o archivos ZIP Shapefile, SHP, GeoJSON, KML, CSV y TXT. La lógica identifica ámbito, distrito, zonificación, parámetros aplicables, proximidad a vías y afectaciones referenciales por alta tensión o ferrocarril.")
add_body(doc, "La proximidad vial se evalúa con un umbral de 25 metros. Los cálculos de área, distancia y perímetro se realizan en UTM 18S para conservar unidades métricas.")

doc.add_heading("Reporte informativo", level=2)
add_body(doc, "El reporte reúne identificación de la consulta, coordenadas, ámbito, distrito, zonificación, parámetros, sistema vial y afectaciones. El motor principal utiliza Quarto y Typst; cuando el entorno lo permite, pagedown y Chrome actúan como alternativa. El documento mantiene una advertencia visible sobre su carácter exclusivamente informativo.")

doc.add_heading("Métricas de uso", level=2)
add_body(doc, "La base de datos mantiene un contador de visitas y otro de constancias generadas. La visita se deduplica por navegador y día. El identificador aleatorio del navegador se transforma a SHA 256 en la base de datos; no se almacena dirección IP ni información personal. El contador de constancias se incrementa después de completar la generación del archivo.")

doc.add_heading("Seguridad y protección de datos", level=1)
add_table(doc, ["Control", "Implementación", "Finalidad"], [
    ("Secretos", "Variables de entorno y exclusión de archivos .env y .Renviron", "Evitar credenciales en el repositorio"),
    ("Transporte", "Conexión PostgreSQL con sslmode require", "Proteger datos en tránsito"),
    ("Acceso espacial", "RLS y revocación para anon y authenticated", "Impedir descarga directa de las capas"),
    ("Funciones públicas", "RPC específicas para métricas", "Exponer solo operaciones necesarias"),
    ("Carga", "Confirmación de project ref y cadena REEMPLAZAR CAPAS PDU", "Reducir el riesgo de cargar en otro proyecto"),
    ("Transacción", "Reemplazo completo dentro de dbWithTransaction", "Evitar conjuntos parciales si ocurre un error"),
    ("Privacidad", "Hash del identificador de visita y ausencia de IP", "Minimizar datos de seguimiento"),
], widths=[3.2, 7.0, 6.3], small=True)
add_body(doc, "La tabla histórica de reportes ciudadanos se conserva para no perder información previa, pero la política de inserción anónima se elimina. La interfaz pública tampoco ofrece formularios de aportes u observaciones.")

doc.add_heading("Despliegue y operación", level=1)
add_body(doc, "Supabase almacena la información espacial y ofrece PostgreSQL PostGIS y funciones RPC. Posit Connect ejecuta el proceso R Shiny. Esta separación debe mantenerse: Supabase no sustituye el entorno de ejecución de R y Posit Connect no debe contener copias de las capas oficiales.")
add_table(doc, ["Variable", "Obligatoria", "Uso"], [
    ("SUPABASE PROJECT REF", "Sí para carga", "Verifica el proyecto institucional de destino"),
    ("SUPABASE HOST", "Sí", "Host directo o pooler de PostgreSQL"),
    ("SUPABASE PORT", "Sí", "Puerto 5432 o el asignado por el pooler"),
    ("SUPABASE DB", "Sí", "Base de datos PostgreSQL"),
    ("SUPABASE USER", "Sí", "Usuario de conexión institucional"),
    ("SUPABASE PASSWORD", "Sí", "Contraseña almacenada como secreto"),
    ("SUPABASE IMPORT CONFIRM", "Solo carga", "Autoriza el reemplazo de las capas previstas"),
    ("SUPABASE URL y ANON KEY", "Opcionales", "Acceso REST para métricas y RPC"),
], widths=[4.6, 2.5, 9.4], small=True)

doc.add_heading("Validación realizada", level=1)
add_body(doc, "La implementación se verificó mediante análisis de sintaxis R, ejecución de la suite testthat, revisión del diff y pruebas específicas de la nueva fuente espacial. Los resultados disponibles son los siguientes.")
add_table(doc, ["Verificación", "Resultado", "Evidencia"], [
    ("Sintaxis de app y módulos", "Conforme", "Los archivos R analizados no presentan errores de parseo"),
    ("Pruebas espaciales existentes", "Conforme", "Once expectativas aprobadas"),
    ("Pruebas de integración espacial", "Conforme", "Se exige PostgreSQL, se valida el catálogo de tablas y la paleta normativa"),
    ("Diff de Git", "Conforme", "Sin errores de espacios o marcadores inválidos"),
    ("Manifiesto de Posit", "Conforme", "Los dos RDS fueron retirados de los archivos de despliegue"),
    ("Carga remota en Supabase", "Pendiente", "Requiere host, usuario y contraseña institucionales"),
    ("Prueba de extremo a extremo", "Pendiente", "Debe ejecutarse después de poblar el esquema public"),
], widths=[4.4, 2.6, 9.5])

doc.add_heading("Estado de implementación", level=1)
add_table(doc, ["Elemento", "Estado", "Observación"], [
    ("Código Shiny con lectura PostGIS", "Completado", "Supabase es la fuente espacial obligatoria"),
    ("Migraciones SQL", "Completado", "PostGIS, auditoría, métricas y cierre del buzón"),
    ("Carga transaccional", "Completado", "Incluye índices GiST y registro de importaciones"),
    ("Contadores públicos", "Completado en código", "Requieren aplicar la migración y desplegar"),
    ("Retiro de aportes", "Completado en interfaz y base", "Conviene eliminar el código residual oculto"),
    ("Población del proyecto Supabase", "Pendiente", "La captura revisada muestra el esquema public sin tablas"),
    ("Secretos en Posit Connect", "Pendiente", "Deben configurarse antes de iniciar la aplicación"),
], widths=[4.6, 3.2, 8.7])

doc.add_page_break()
doc.add_heading("Evaluación de la interfaz", level=1)
add_body(doc, "La interfaz utiliza una paleta institucional de verde, dorado y neutros, incorpora el escudo municipal, el logotipo del PDU y una navegación basada en bslib. El mapa Leaflet ocupa el área principal y el panel lateral concentra la selección de modalidad, coordenadas, archivos, resultados y descargas.")

doc.add_heading("Aspectos favorables", level=2)
add_bullet(doc, "La identidad visual permite reconocer a la Municipalidad y el PDU desde la cabecera.")
add_bullet(doc, "Los modos de consulta por punto y polígono responden a situaciones reales de uso ciudadano y técnico.")
add_bullet(doc, "La ficha muestra ámbito, zonificación, vías y parámetros sin exigir conocimientos especializados de SIG.")
add_bullet(doc, "El botón de GPS facilita la consulta desde teléfonos en campo.")
add_bullet(doc, "Los estados de foco en formularios y la separación visual de resultados mejoran la lectura.")
add_bullet(doc, "El reporte conserva el aviso sobre su carácter informativo y ofrece continuidad entre consulta y descarga.")

doc.add_heading("Problemas observados", level=2)
add_table(doc, ["Hallazgo", "Impacto", "Prioridad"], [
    ("La cabecera reúne nombre extenso, fecha, estado de nube y dos contadores", "Puede saturarse en anchos intermedios y distraer de la consulta", "Alta"),
    ("El panel lateral tiene ancho fijo de 410 píxeles y muchos controles", "Reduce el mapa y obliga a desplazamiento, especialmente en portátiles y móviles", "Alta"),
    ("Se combinan gradientes, sombras, badges, bordes y movimientos hover", "La interfaz se percibe más promocional que administrativa", "Media"),
    ("La leyenda de zonificación puede incluir numerosas subzonas", "Dificulta encontrar una categoría y ocupa espacio del mapa", "Alta"),
    ("Los colores son el principal diferenciador de zonas y alertas", "Usuarios con deficiencias de visión cromática pueden perder información", "Alta"),
    ("El texto Cloud DB Activa expone terminología técnica", "No comunica al ciudadano si los datos están vigentes", "Media"),
    ("La denominación alterna entre reporte y constancia", "Puede generar confusión sobre el valor administrativo del archivo", "Alta"),
    ("El buzón está oculto con una condición falsa pero conserva código", "Aumenta el mantenimiento y el riesgo de reactivación accidental", "Media"),
    ("No existe un recorrido introductorio ni estado inicial orientador", "Una persona nueva puede no saber qué modalidad elegir", "Media"),
    ("La altura del mapa depende de un cálculo fijo respecto de la cabecera", "Puede producir cortes cuando cambia el tamaño del navbar", "Media"),
], widths=[7.4, 6.8, 2.3], small=True)

doc.add_heading("Revisión de la salida cartográfica", level=2)
map_path = ROOT / "test_map.png"
if map_path.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = p.add_run().add_picture(str(map_path), width=Cm(15.2))
    set_picture_alt(pic, "Mapa de prueba con predio resaltado en rojo y vías propuestas en colores púrpura y verde")
    add_caption(doc, "Figura 2 Muestra de verificación cartográfica disponible en el repositorio")
add_body(doc, "La muestra confirma que el predio consultado se distingue mediante borde rojo y relleno semitransparente, mientras las vías conservan colores por jerarquía. Como mejora, el mapa debe incluir rótulo de escala, norte, fuente, fecha de actualización y una leyenda que identifique también la zonificación visible. El fondo gris debe reemplazarse o explicarse cuando represente categorías sin color asignado.")

doc.add_heading("Propuesta de mejora de interfaz", level=1)
doc.add_heading("Prioridad P0 antes de publicación", level=2)
add_table(doc, ["Mejora", "Acción propuesta", "Criterio de aceptación"], [
    ("Fuente institucional", "Cargar Supabase y mostrar fecha de última importación", "La aplicación inicia solo con las nueve tablas completas y muestra la fecha de actualización"),
    ("Lenguaje legal", "Usar Reporte informativo en toda la aplicación", "Ningún botón o título denomina certificado oficial al archivo"),
    ("Código residual", "Eliminar completamente el nav panel y observadores del buzón", "No quedan controles ni funciones públicas de envío ciudadano"),
    ("Estado del servicio", "Cambiar Cloud DB Activa por Datos territoriales actualizados", "El mensaje es comprensible sin terminología de infraestructura"),
    ("Errores", "Diseñar mensajes para conexión, tabla vacía y PDF no disponible", "Cada fallo indica qué ocurrió y qué puede hacer la persona usuaria"),
    ("Privacidad", "Añadir una nota breve sobre métricas", "La página explica que no se almacena IP ni datos personales"),
], widths=[3.5, 7.2, 5.8], small=True)

doc.add_heading("Prioridad P1 experiencia de consulta", level=2)
add_table(doc, ["Mejora", "Diseño recomendado", "Beneficio"], [
    ("Flujo guiado", "Organizar la consulta en tres pasos: ubicación, resultado y reporte", "Reduce decisiones simultáneas"),
    ("Cabecera compacta", "Mantener escudo, nombre corto y estado de actualización; mover métricas a Información", "Libera espacio y conserva el carácter institucional"),
    ("Panel adaptable", "Usar panel plegable en escritorio y hoja inferior en móvil", "Amplía el mapa y mejora el uso táctil"),
    ("Leyenda", "Agrupar por macrozona, permitir búsqueda y mostrar solo capas activas", "Facilita la interpretación del mapa"),
    ("Resultado principal", "Mostrar primero zona, condición de ámbito y alertas; dejar detalles en acordeones", "Prioriza la información que responde la consulta"),
    ("Accesibilidad", "Añadir patrones, etiquetas, ARIA, orden de foco y contraste WCAG 2.2 AA", "Evita depender únicamente del color"),
    ("Primera visita", "Incluir una explicación de 30 segundos y ejemplos de coordenadas", "Reduce errores de entrada"),
], widths=[3.4, 7.8, 5.3], small=True)

doc.add_heading("Prioridad P2 evolución operativa", level=2)
add_bullet(doc, "Consultar por extensión visible o usar teselas vectoriales si el volumen espacial aumenta.")
add_bullet(doc, "Crear un tablero interno de consultas, importaciones, errores y constancias, separado de la interfaz pública.")
add_bullet(doc, "Automatizar la actualización de capas con validaciones de geometría, esquema, conteos y fecha de vigencia.")
add_bullet(doc, "Incorporar monitoreo de disponibilidad y tiempos de respuesta de Supabase y Posit Connect.")
add_bullet(doc, "Realizar pruebas de usabilidad con población, personal de catastro, desarrollo urbano y atención al ciudadano.")

doc.add_heading("Lineamientos visuales institucionales", level=2)
add_body(doc, "Se recomienda una presentación sobria basada en fondo blanco, verde institucional para acciones principales y dorado solo como acento. Deben reducirse los gradientes, sombras y animaciones. Los títulos deben usar lenguaje administrativo directo y los estados deben expresar el significado para el ciudadano, no el nombre de la tecnología.")
add_table(doc, ["Elemento", "Lineamiento"], [
    ("Cabecera", "Una sola línea principal con identidad y una segunda línea breve para vigencia de datos"),
    ("Tipografía", "Inter o Arial con tamaños mínimos de 16 píxeles para controles y 14 para texto auxiliar"),
    ("Botones", "Una acción primaria por etapa; acciones secundarias con borde y sin gradiente"),
    ("Colores", "Verde para acción, ámbar para advertencia y rojo solo para riesgo o error"),
    ("Tarjetas", "Bordes suaves y sombra mínima; evitar movimiento al pasar el cursor"),
    ("Mapas", "Leyenda plegable, fuente, fecha, escala, norte y etiquetas legibles"),
    ("Móvil", "Controles de al menos 44 por 44 píxeles y panel deslizable"),
], widths=[3.6, 12.9])

doc.add_heading("Hoja de ruta", level=1)
add_table(doc, ["Fase", "Actividades", "Responsable propuesto", "Producto verificable"], [
    ("1 Activación de datos", "Configurar secretos, aplicar migraciones, cargar capas y verificar conteos", "Administrador Supabase y equipo SIG", "Esquema public poblado y script 05 conforme"),
    ("2 Publicación controlada", "Configurar Posit, desplegar y ejecutar pruebas funcionales", "Equipo de aplicación", "GeoVisor accesible con lectura PostGIS"),
    ("3 Mejora P0", "Unificar textos, retirar código residual, fecha de datos y mensajes de error", "Desarrollo y responsable funcional", "Versión candidata institucional"),
    ("4 Mejora P1", "Rediseño responsive, accesibilidad, leyenda y flujo guiado", "UX desarrollo y comunicaciones", "Prueba de usabilidad aprobada"),
    ("5 Operación", "Monitoreo, respaldos, actualización de capas y revisión trimestral", "Unidad responsable", "Registro periódico de operación"),
], widths=[2.7, 6.2, 3.8, 3.8], small=True)

doc.add_heading("Riesgos y medidas de control", level=1)
add_table(doc, ["Riesgo", "Consecuencia", "Control recomendado"], [
    ("Credenciales no configuradas", "La aplicación no inicia", "Configurar secretos en Posit y probar con el script de verificación"),
    ("Carga parcial o fallida", "Datos inconsistentes", "Mantener la transacción y respaldar antes de cada actualización"),
    ("Cambio de columnas", "Errores en consultas o fichas", "Validar esquema y ejecutar pruebas automáticas antes de publicar"),
    ("Uso del reporte como documento oficial", "Interpretación administrativa incorrecta", "Advertencia visible y denominación uniforme Reporte informativo"),
    ("Saturación del mapa", "Tiempos de carga altos", "Simplificación controlada, consulta por extensión o teselas"),
    ("Dependencia del color", "Barreras de accesibilidad", "Patrones, iconos, texto y contraste verificado"),
    ("Reactivación accidental del buzón", "Recepción de datos personales sin proceso de atención", "Eliminar código y conservar la revocación de inserción anónima"),
], widths=[4.2, 5.2, 7.1], small=True)

doc.add_heading("Conclusiones y recomendaciones", level=1)
add_body(doc, "El GeoVisor cuenta con las funciones necesarias para informar a la población sobre la situación territorial de un punto o predio respecto del PDU de Anta. La transición técnica establece Supabase PostGIS como fuente espacial única y evita que la aplicación publique copias locales desactualizadas.")
add_body(doc, "La implementación de código está validada, pero el producto no debe considerarse operativo hasta completar la carga en el proyecto Supabase, configurar los secretos en Posit Connect y ejecutar una prueba de extremo a extremo. La captura del panel revisada muestra que el esquema public todavía no contiene tablas.")
add_body(doc, "Antes de la publicación institucional se recomienda ejecutar todas las acciones P0, realizar una revisión funcional con la Gerencia de Desarrollo Urbano y Rural y validar el lenguaje del reporte informativo. La siguiente etapa debe concentrarse en accesibilidad, experiencia móvil y simplificación visual.")

doc.add_page_break()
doc.add_heading("Anexo A Procedimiento de activación", level=1)
add_manual_number(doc, 1, "Crear un archivo .Renviron local o configurar secretos en Posit Connect con host, puerto, base, usuario y contraseña de Supabase.")
add_manual_number(doc, 2, "Ejecutar scripts 03 aplicar migraciones Supabase R desde la raíz del proyecto.")
add_manual_number(doc, 3, "Ejecutar scripts 04 cargar Supabase R con la confirmación REEMPLAZAR CAPAS PDU.")
add_manual_number(doc, 4, "Ejecutar scripts 05 verificar Supabase R y conservar el resultado como evidencia de despliegue.")
add_manual_number(doc, 5, "Publicar la aplicación en Posit Connect y probar consulta puntual, consulta por polígono, catálogo, reporte y contadores.")
add_manual_number(doc, 6, "Registrar la fecha de importación y la versión del repositorio utilizada.")

doc.add_heading("Anexo B Inventario de archivos técnicos", level=1)
add_table(doc, ["Archivo", "Finalidad"], [
    ("app.R", "Interfaz Shiny, análisis de consultas y generación de reportes"),
    ("R supabase utils R", "Conexión, carga PostGIS, métricas, auditoría y avisos"),
    ("R spatial utils R", "Validación y normalización de geometrías prediales"),
    ("supabase migrations 001 postgis pdu sql", "PostGIS y trazabilidad de importaciones"),
    ("supabase migrations 002 metricas geovisor sql", "Contadores, RPC y cierre de la inserción ciudadana"),
    ("scripts 03 aplicar migraciones Supabase R", "Aplicación idempotente de migraciones"),
    ("scripts 04 cargar Supabase R", "Carga transaccional e índices espaciales"),
    ("scripts 05 verificar Supabase R", "Prueba previa al despliegue"),
    ("constancia pdu template qmd", "Plantilla del reporte informativo"),
    ("www custom css", "Sistema visual de la aplicación"),
    ("manifest json", "Dependencias y archivos del despliegue Posit"),
], widths=[7.0, 9.5])

doc.add_page_break()
annex_c_heading = doc.add_heading("Anexo C Criterios de aceptación operativa", level=1)
annex_c_heading.paragraph_format.left_indent = Cm(1.0)
annex_c_heading.paragraph_format.space_before = Pt(28)
criteria = [
    "Las nueve tablas requeridas existen y contienen los conteos esperados o una actualización aprobada.",
    "La aplicación inicia sin archivos RDS en el paquete de Posit Connect.",
    "Una consulta puntual válida identifica ámbito distrito y zonificación.",
    "Una consulta por polígono calcula área perímetro porcentaje de ámbito y superposición de zonas.",
    "El reporte informativo se genera y el contador de constancias aumenta una vez.",
    "Una visita repetida en el mismo navegador y día no incrementa el contador nuevamente.",
    "No existe una opción pública para enviar aportes y el rol anónimo no puede insertar reportes ciudadanos.",
    "Los mensajes legales distinguen el reporte informativo del certificado oficial municipal.",
    "La interfaz es usable con teclado y en una pantalla móvil de ancho reducido.",
    "La fecha y la fuente de actualización de los datos son visibles para la población.",
]
for item in criteria:
    add_bullet(doc, item)

# Document properties and final save
doc.core_properties.title = "Informe de implementación del GeoVisor PDU Anta 2024 2034"
doc.core_properties.subject = "Arquitectura Supabase PostGIS aplicación Shiny y evaluación de interfaz"
doc.core_properties.author = "Municipalidad Provincial de Anta"
doc.core_properties.keywords = "PDU Anta GeoVisor Shiny Supabase PostGIS gestión territorial"
doc.core_properties.comments = "Documento técnico institucional"

settings = doc.settings._element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

doc.save(OUTPUT)
print(OUTPUT)
