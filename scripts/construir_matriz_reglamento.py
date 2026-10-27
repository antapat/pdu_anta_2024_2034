"""Extrae el reglamento final y vincula sus categorías con el DBF original.

Ejecutar con Python y python-docx. No modifica el DOCX ni el Shapefile.
Los valores se transcriben como texto; no se sustituyen por la matriz antigua.
"""
from pathlib import Path
from collections import Counter
from docx import Document
import csv
import hashlib
import json
import re
import struct
import unicodedata

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "TOMO III_REGLAMENTO_2024-2034.docx"
OUT = ROOT / "outputs/reglamento"
NORM = ROOT / "normativa"
FIELDS = {
    "Densidad neta": "DENSIDAD_NETA_MAX",
    "Área mínima": "LOTE_MINIMO",
    "Frente mínimo": "FRENTE_MINIMO",
    "Altura máxima": "ALTURA_MAXIMA",
    "Retiros": "RETIRO_FRONTAL",
    "Coeficiente": "COEFICIENTE_EDIF",
    "Porcentaje mínimo": "AREA_LIBRE_MIN",
    "Número mínimo": "ESTACIONAMIENTOS",
    "Volad": "VOLADIZOS",
    "Usos de suelo": "USOS_PERMITIDOS",
}
NA = "No establecido en este apartado del reglamento"


def norm(text):
    return re.sub(r"[^A-Z0-9]", "", unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().upper())


def clean(text):
    return "\n".join(re.sub(r"[ \t]+", " ", p).strip() for p in text.splitlines()).strip()


def write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_dbf(path):
    raw = path.read_bytes()
    count = struct.unpack_from("<I", raw, 4)[0]
    header, size = struct.unpack_from("<HH", raw, 8)
    fields = []
    pos = 32
    while raw[pos] != 13:
        desc = raw[pos:pos + 32]
        fields.append((desc[:11].split(b"\0")[0].decode("ascii"), desc[16]))
        pos += 32
    records = []
    for i in range(count):
        record = raw[header + i * size:header + (i + 1) * size]
        if record[0:1] == b"*":
            raise ValueError("DBF con registros borrados: requiere conciliación con SHP")
        row = {"FILA_SHP": i + 1}
        pos = 1
        for field, length in fields:
            row[field] = record[pos:pos + length].decode("utf-8").strip().strip("\0")
            pos += length
        records.append(row)
    return records


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    NORM.mkdir(exist_ok=True)
    doc = Document(SOURCE)
    paragraphs = {i + 1: clean(p.text) for i, p in enumerate(doc.paragraphs)}
    tables = doc.tables
    hash_doc = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    # Ubicación exacta: índice de tabla XML y párrafo del cuerpo, sin inventar páginas.
    contexts = []
    pi = ti = 0
    for element in doc.element.body:
        if element.tag.endswith("}p"):
            pi += 1
        elif element.tag.endswith("}tbl"):
            ti += 1
            contexts.append({"xml": ti, "antes": pi})
    catalog = []
    for row in tables[0].rows:
        c = [clean(cell.text) for cell in row.cells]
        code = c[2].replace("’", "'")
        macro = re.search(r"\(([^()]+)\)", c[1]).group(1)
        catalog.append({"COD_S_ZONA": code, "COD_ZONA": macro, "NOMBRE_COMPLETO": c[3]})
    catalog.append({"COD_S_ZONA": "ZRE-CBIP", "COD_ZONA": "ZRE",
                    "NOMBRE_COMPLETO": "Zona de Reglamentación Especial por Compromiso con Bien Inmueble Prehispánico"})
    assert len(catalog) == 40 and len({x['COD_S_ZONA'] for x in catalog}) == 40
    residential = [c["COD_S_ZONA"] for c in catalog[:18]]
    table_for = {c: (i + 2, i + 15) for i, c in enumerate(residential)}
    table_for.update({"E1": (22, 34), "E1'": (22, 34), "E2": (23, 35),
                      "H1": (24, 36), "H2": (25, 37), "H3": (26, 38),
                      "ZRP": (27, 39), "ZOU": (28, 40), "I1": (29, 41), "I2": (30, 42)})
    spans = {"ZA": (360, 371), "ZM": (373, 408), "ZPA-RH": (418, 428),
             "ZPA-CE": (431, 440), "ZPA-IERE": (441, 445),
             "ZRE-NU-P": (488, 504), "ZRE-NU-RE": (505, 522),
             "ZRE-PMA-S": (543, 560), "ZRE-PA-C": (561, 563),
             "ZRE-PA-HYI": (564, 571), "ZRE-VHC-CH": (572, 580), "ZRE-CBIP": (581, 584)}
    allowed = {"ZA": [362, 367], "ZM": [], "ZPA-RH": [424, 425, 426, 427, 428],
               "ZPA-CE": [437, 438, 439, 440], "ZPA-IERE": [444, 445],
               "ZRE-NU-P": [500, 504], "ZRE-NU-RE": [517, 518],
               "ZRE-PMA-S": [556], "ZRE-PA-C": [], "ZRE-PA-HYI": [],
               "ZRE-VHC-CH": [580], "ZRE-CBIP": [584]}
    restricted = {"ZA": [362], "ZM": [375, 377], "ZPA-RH": [419, 421, 422],
                  "ZPA-CE": [432, 435, 436], "ZPA-IERE": [442],
                  "ZRE-NU-P": [502, 503, 504], "ZRE-NU-RE": [521, 522],
                  "ZRE-PMA-S": [545, 558, 559, 560], "ZRE-PA-C": [563],
                  "ZRE-PA-HYI": [566, 567, 568], "ZRE-VHC-CH": [575, 580], "ZRE-CBIP": [584]}
    joinp = lambda ids: "\n".join(f"[P{i}] {paragraphs[i]}" for i in ids if paragraphs.get(i))
    matrix, evidence = [], []
    for item in catalog:
        code = item["COD_S_ZONA"]
        row = dict(item)
        row.update({key: NA for key in FIELDS.values()})
        row.update({"RESTRICCIONES": "Sin restricciones adicionales en la tabla individual; revisar notas.",
                    "NOTAS_APLICACION": "", "FUENTE_NORMATIVA": SOURCE.name,
                    "SHA256_DOCUMENTO": hash_doc,
                    "ESTADO_VALIDACION": "TRANSCRITO DEL REGLAMENTO FINAL",
                    "OBSERVACIONES_FUENTE": ""})
        if code in table_for:
            xml, number = table_for[code]
            row["TABLA_FUENTE"] = f"Tabla N° {number} (tabla XML {xml})"
            before = contexts[xml - 1]["antes"]
            row["REFERENCIA_DOCUMENTAL"] = f"P{before}; {row['TABLA_FUENTE']}"
            rows = tables[xml - 1].rows
            restriction_mode = False
            restrictions = []
            for ri, table_row in enumerate(rows):
                cells = list(dict.fromkeys(clean(c.text) for c in table_row.cells))
                label = cells[0]
                if norm(label) == "RESTRICCIONES":
                    restriction_mode = True
                    if len(cells) > 1:
                        restrictions.append(cells[-1])
                    continue
                if restriction_mode:
                    restrictions.append("\n".join(cells))
                    continue
                field = next((dest for prefix, dest in FIELDS.items() if norm(label).startswith(norm(prefix))), None)
                if field:
                    value = clean(table_row.cells[-1].text)
                    row[field] = value or "Celda vacía en la fuente"
                    evidence.append({"COD_S_ZONA": code, "CAMPO": field, "VALOR_LITERAL": value,
                                     "REFERENCIA": f"Tabla {number}; XML {xml}; fila {ri + 1}",
                                     "ETIQUETA_ORIGINAL": label})
            assert all(row[k] != NA for k in FIELDS.values()), (code, row)
            if restrictions:
                row["RESTRICCIONES"] = "\n".join(dict.fromkeys(restrictions))
            after = contexts[xml]["antes"] if xml < len(contexts) else 743
            # Notas numeradas consecutivas antes del siguiente encabezado.
            note_ids = []
            started = False
            for p in range(before + 1, min(after + 1, before + 15)):
                text = paragraphs.get(p, "")
                if re.match(r"^\(\d+\)", text):
                    started = True
                    note_ids.append(p)
                elif started and text and not text.startswith("Fuente"):
                    break
            row["NOTAS_APLICACION"] = joinp(note_ids)
            if code == "H3":
                row["OBSERVACIONES_FUENTE"] = "La leyenda de Tabla 38 y P314–318 identifican Hospital General H3; el encabezado interno dice Posta Médica H1. Se transcriben los valores de Tabla 38, sin sustituirlos por H1."
            if code == "I2":
                row["OBSERVACIONES_FUENTE"] = "La leyenda de Tabla 42 identifica Industria Liviana I2; el encabezado interno dice Industria Elemental y Complementaria I2. Se conserva la categoría de la leyenda."
            if code == "ZDMB-CR":
                row["OBSERVACIONES_FUENTE"] = "Tabla 32 incluye Taller entre usos; Tabla 33 resume usos sin Taller. Se muestran los valores de la tabla individual y esta discrepancia, sin resolverla administrativamente."
        else:
            lo, hi = spans[code]
            row["TABLA_FUENTE"] = "Tabla N° 44 y disposiciones de ZM" if code == "ZM" else "Disposiciones específicas; sin tabla numérica individual"
            row["REFERENCIA_DOCUMENTAL"] = f"P{lo}–P{hi}" + ("; Tabla 43, columna ZA" if code == "ZA" else "")
            row["USOS_PERMITIDOS"] = joinp(allowed[code]) or "Sujeto a las disposiciones y al instrumento específico indicado en restricciones."
            row["RESTRICCIONES"] = joinp(restricted[code])
            row["NOTAS_APLICACION"] = joinp([i for i in range(lo, hi + 1) if i not in allowed[code] + restricted[code] and paragraphs.get(i) and len(paragraphs[i]) > 100])
            if code == "ZM":
                row["USOS_PERMITIDOS"] = "\n".join(" | ".join(dict.fromkeys(clean(c.text) for c in r.cells)) for r in tables[31].rows[1:])
            if code in ("ZPA-CE", "ZPA-IERE"):
                row["NOTAS_APLICACION"] += "\n" + joinp(range(446, 466))
                row["REFERENCIA_DOCUMENTAL"] += "; disposiciones compartidas P446–P465"
            if code == "ZRE-CBIP":
                row["OBSERVACIONES_FUENTE"] = "Categoría desarrollada en P581–584 y presente en la cartografía; no incluida en Tabla 14."
            if code == "ZRE-VHC-CH":
                row["ALTURA_MAXIMA"] = "Consultar parámetros de OM N° 017–2010-MPA-A/SG; el documento remite a ellos y no los reproduce."
        matrix.append(row)

    plans = {"ZRE-PMA-S": (592, 609), "ZRE-PA-C": (610, 617),
             "ZRE-PA-HYI": (618, 622), "ZRE-VHC-CH": (623, 633), "ZRE-CBIP": (634, 640)}
    for row in matrix:
        code = row['COD_S_ZONA']
        if code in plans:
            lo, hi = plans[code]
            row['NOTAS_APLICACION'] += '\n' + joinp(range(lo, hi + 1))
            row['REFERENCIA_DOCUMENTAL'] += f'; planes P{lo}–P{hi}'
        if code == 'ZA':
            row['NOTAS_APLICACION'] += '\n[Tabla 43, columna ZA] ' + clean(tables[30].rows[2].cells[-1].text)
        row['CONDICIONES_GENERALES'] = (
            'Compatibilidad de actividades: el reglamento exige consultar el IUUAU '
            '(P643–674; Anexo 2). Esta matriz no reproduce el cuadro de actividades del anexo '
            'ni determina la compatibilidad de una actividad económica concreta. '
            'SLO: las disposiciones P677–692 vinculan las alturas con las restricciones '
            'aeronáuticas; el Shapefile de zonificación no permite verificar por sí solo '
            'la afectación espacial por SLO. '
            'Predios colindantes al ferrocarril: aplicar la remisión de P261 al artículo 19 '
            'del D.S. 032-2005-MTC. La consulta cartográfica no sustituye las verificaciones sectoriales.'
        )

    # Correspondencia explícita por descripción; SIMB_ZONI se conserva como evidencia.
    by_desc = {norm(r['NOMBRE_COMPLETO']): r['COD_S_ZONA'] for r in matrix}
    aliases = {
        "INDUSTRIA ELEMENTAL": "I1",
        "ZONA DE PROTECCION AMBIENTAL POR INTERVENCION ESPECIAL Y RECUPERACION DE ECOS": "ZPA-IERE",
        "ZONA DE REGLAMENTACION ESPECIAL NO URBANIZABLE POR PELIGRO ALTO Y MUY ALTO": "ZRE-NU-P",
        "ZONA DE REGLAMENTACION ESPECIAL NO URBANIZABLE POR RIESGO ELECTRICO": "ZRE-NU-RE",
        "ZONA DE REGLAMENTACION ESPECIAL POR PROTECCION AMBIENTAL HUMEDAL YUNGAQUI-": "ZRE-PA-HYI",
        "ZONA DE REGLAMENTACION ESPECIAL POR VALORES HISTORICOS CULTURALES - CENTRO": "ZRE-VHC-CH",
    }
    by_alias = {norm(k): v for k, v in aliases.items()}
    features = read_dbf(ROOT / "data/P_080301_Zonificacion_urbana.dbf")
    groups = Counter((r['SIMB_ZONI'], r['SCLAS_ZONI'], r['NOM_ZONI']) for r in features)
    lookup = {r['COD_S_ZONA']: r for r in matrix}
    crosswalk = []
    for (symbol, desc, name), count in sorted(groups.items()):
        code = by_desc.get(norm(desc)) or by_alias.get(norm(desc))
        if not code:
            raise ValueError(f"Descripción sin equivalencia explícita: {desc}")
        expected = lookup[code]['COD_ZONA']
        normalized_symbol = re.sub(r"-R(I|II)?$", "", symbol)
        warning = "" if normalized_symbol == expected else f"Símbolo original {symbol} no coincide con macrozona normativa {expected}; vínculo por SCLAS_ZONI. No se modifica el Shapefile."
        crosswalk.append({"SIMB_ZONI": symbol, "SCLAS_ZONI": desc, "NOM_ZONI": name,
                          "COD_S_ZONA": code, "COD_ZONA": expected, "POLIGONOS": count,
                          "METODO_VINCULO": "Descripción normalizada exacta" if norm(desc) in by_desc else "Alias explícito documentado",
                          "OBSERVACION_VINCULO": warning})
    key = {(r['SIMB_ZONI'], r['SCLAS_ZONI'], r['NOM_ZONI']): r for r in crosswalk}
    counts = Counter()
    per_feature = []
    for feature in features:
        link = key[(feature['SIMB_ZONI'], feature['SCLAS_ZONI'], feature['NOM_ZONI'])]
        counts[link['COD_S_ZONA']] += 1
        per_feature.append({k: feature[k] for k in ['FILA_SHP', 'OBJECTID', 'IDLOTE', 'ID_MANZANA', 'SIMB_ZONI', 'SCLAS_ZONI']} |
                           {k: link[k] for k in ['COD_S_ZONA', 'COD_ZONA', 'METODO_VINCULO', 'OBSERVACION_VINCULO']})
    for row in matrix:
        row['POLIGONOS_VINCULADOS'] = counts[row['COD_S_ZONA']]
        warnings = [r['OBSERVACION_VINCULO'] + f" ({r['POLIGONOS']} polígonos)." for r in crosswalk if r['COD_S_ZONA'] == row['COD_S_ZONA'] and r['OBSERVACION_VINCULO']]
        row['OBSERVACIONES_CARTOGRAFIA'] = "\n".join(warnings)
    write_csv(NORM / 'matriz_reglamento_final.csv', matrix)
    write_csv(NORM / 'vinculo_zonificacion_reglamento.csv', crosswalk)
    write_csv(OUT / 'vinculo_por_poligono.csv', per_feature)
    write_csv(OUT / 'evidencia_parametros.csv', evidence)
    payload = {"matriz": matrix, "vinculo": crosswalk, "evidencia": evidence,
               "parrafos": paragraphs, "fuente": SOURCE.name, "sha256": hash_doc,
               "poligonos": len(features)}
    (OUT / 'matriz_revision.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Matriz: {len(matrix)} categorías; {len(crosswalk)} combinaciones cartográficas; {len(features)} polígonos vinculados")
    print("Observaciones cartográficas:", [(r['COD_S_ZONA'], r['POLIGONOS']) for r in crosswalk if r['OBSERVACION_VINCULO']])


if __name__ == '__main__':
    main()
